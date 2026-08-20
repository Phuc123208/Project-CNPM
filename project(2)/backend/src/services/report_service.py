"""Module 6 - Reporting. Generates traffic-analysis / forecast / experiment
reports as downloadable PDF or Excel files."""
import os
import uuid
from datetime import datetime
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from domain.exceptions import NotFoundException
from infrastructure.repositories.report_repository import ReportRepository
from infrastructure.repositories.experiment_repository import ExperimentRepository
from infrastructure.repositories.dataset_repository import DatasetRepository
from services.analysis_service import AnalysisService


class ReportService:
    def __init__(self, report_folder="static/reports"):
        self.report_folder = report_folder
        os.makedirs(self.report_folder, exist_ok=True)
        self.repository = ReportRepository()
        self.experiment_repo = ExperimentRepository()
        self.dataset_repo = DatasetRepository()
        self.analysis_service = AnalysisService(self.dataset_repo)

    def _new_path(self, prefix, ext):
        filename = f"{prefix}_{uuid.uuid4().hex[:8]}.{ext}"
        return os.path.join(self.report_folder, filename), filename

    # ---------------- Traffic analysis report ----------------
    def generate_traffic_report(self, user_id, version_id, fmt="pdf"):
        kpis = self.analysis_service.get_kpis(version_id)
        hotspots = self.analysis_service.get_hotspots(version_id)
        peak_hours = self.analysis_service.get_peak_hours(version_id)

        if fmt == "excel":
            path, filename = self._new_path("traffic_report", "xlsx")
            wb = Workbook()
            ws = wb.active
            ws.title = "KPIs"
            ws.append(["Metric", "Value"])
            for k, v in kpis.items():
                ws.append([k, v])
            ws2 = wb.create_sheet("Hotspots")
            ws2.append(["segment_id", "avg_density", "avg_speed", "total_vehicle_count"])
            for h in hotspots:
                ws2.append([h.get("segment_id"), h.get("avg_density"), h.get("avg_speed"),
                            h.get("total_vehicle_count")])
            ws3 = wb.create_sheet("Peak Hours")
            ws3.append(["hour", "avg_vehicle_count", "is_peak"])
            for p in peak_hours:
                ws3.append([p.get("hour"), p.get("avg_vehicle_count"), p.get("is_peak")])
            wb.save(path)
        else:
            fmt = "pdf"
            path, filename = self._new_path("traffic_report", "pdf")
            self._render_pdf(path, "Traffic Analysis Report", [
                ("Key Performance Indicators", [[k, str(v)] for k, v in kpis.items()]),
                ("Congestion Hotspots (Top segments by density)",
                 [["Segment", "Avg Density", "Avg Speed", "Total Vehicles"]] +
                 [[h.get("segment_id"), round(h.get("avg_density", 0), 2),
                   round(h.get("avg_speed", 0), 2), h.get("total_vehicle_count")]
                  for h in hotspots]),
                ("Peak Hour Analysis",
                 [["Hour", "Avg Vehicle Count", "Peak?"]] +
                 [[p.get("hour"), round(p.get("avg_vehicle_count", 0), 2), p.get("is_peak")]
                  for p in peak_hours]),
            ])

        report = self.repository.create(user_id, f"Traffic Analysis Report - v{version_id}",
                                          fmt, filename)
        return report

    # ---------------- Forecast report ----------------
    def generate_forecast_report(self, user_id, experiment_id, fmt="pdf"):
        exp = self.experiment_repo.get(experiment_id)
        if not exp:
            raise NotFoundException("Experiment not found")
        forecasts = self.experiment_repo.get_forecast_results(experiment_id)

        if fmt == "excel":
            path, filename = self._new_path("forecast_report", "xlsx")
            wb = Workbook()
            ws = wb.active
            ws.title = "Forecast"
            ws.append(["forecast_time", "segment_id", "predicted_density", "lower_bound", "upper_bound"])
            for f in forecasts:
                ws.append([str(f.forecast_time), f.segment_id, f.predicted_density,
                           f.lower_bound, f.upper_bound])
            ws2 = wb.create_sheet("Metrics")
            for k, v in (exp.evaluation_metrics or {}).items():
                ws2.append([k, v])
            wb.save(path)
        else:
            fmt = "pdf"
            path, filename = self._new_path("forecast_report", "pdf")
            self._render_pdf(path, f"Forecast Report - {exp.name} ({exp.model_type.upper()})", [
                ("Evaluation Metrics", [[k, str(v)] for k, v in (exp.evaluation_metrics or {}).items()]),
                ("Forecast Results",
                 [["Time", "Segment", "Predicted Density", "Lower", "Upper"]] +
                 [[str(f.forecast_time), f.segment_id, round(f.predicted_density or 0, 2),
                   round(f.lower_bound or 0, 2), round(f.upper_bound or 0, 2)] for f in forecasts]),
            ])

        report = self.repository.create(user_id, f"Forecast Report - {exp.name}", fmt,
                                          filename, experiment_id=experiment_id)
        return report

    # ---------------- Experiment report ----------------
    def generate_experiment_report(self, user_id, experiment_id, fmt="pdf"):
        return self.generate_forecast_report(user_id, experiment_id, fmt)

    def _render_pdf(self, path, title, sections):
        doc = SimpleDocTemplate(path, pagesize=A4)
        styles = getSampleStyleSheet()
        story = [Paragraph(title, styles["Title"]),
                 Paragraph(f"Generated: {datetime.utcnow().isoformat()} UTC", styles["Normal"]),
                 Spacer(1, 16)]
        for heading, rows in sections:
            story.append(Paragraph(heading, styles["Heading2"]))
            story.append(Spacer(1, 6))
            if rows:
                table = Table(rows, hAlign="LEFT")
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2c3e50")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                ]))
                story.append(table)
            story.append(Spacer(1, 16))
        doc.build(story)

    def list_reports(self, user_id=None):
        return self.repository.list(user_id=user_id)

    def get_report(self, report_id):
        r = self.repository.get(report_id)
        if not r:
            raise NotFoundException("Report not found")
        return r
